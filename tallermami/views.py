
import logging
import os
from copy import deepcopy

from django import forms
from django.contrib import admin
from django.contrib import messages
from django.contrib.auth.forms import AdminPasswordChangeForm
from django.contrib.auth import login, logout
from django.contrib.auth.decorators import login_required, user_passes_test
from django.contrib.auth.forms import (
    AuthenticationForm,
    UserChangeForm,
    UserCreationForm,
)
from django.contrib.auth.models import User
from django.core.exceptions import (
    FieldDoesNotExist,
    ImproperlyConfigured,
    ValidationError,
)
from django.core.paginator import Paginator
from django.db import transaction
from django.db.models.deletion import ProtectedError
from django.db.models import OuterRef, Q, Subquery
from django.forms import modelform_factory
from django.http import HttpResponse, HttpResponseNotAllowed
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from .forms import (
    AvanceReparacionForm,
    ClienteForm,
    EnviarWhatsAppForm,
    PresupuestoForm,
    RepuestosPresupuestadosFormSet,
    RegistroForm,
    PresupuestoFinalForm,
    RepuestosUtilizadosFormSet,
    SolicitarTurnoForm,
    TurnoForm,
)
from .models import (
    Cliente,
    Disponibilidad,
    OfertaTurno,
    PerfilUsuario,
    Presupuesto,
    RepuestoUtilizado,
    Trabajador,
    Trabajo,
    Turno,
    Vehiculo,
)
from .services import guardar_turno_con_disponibilidad
from .utils import (
    WhatsAppDeliveryError,
    enviar_correo_taller,
    enviar_whatsapp_a_cliente,
    notificar_avance_reparacion,
    mensaje_avance_whatsapp,
)

logger = logging.getLogger(__name__)

# ==========================================================
# HELPERS DEL EXCEL
# ==========================================================
def obtener_ruta_excel():
    carpeta_excel = os.path.join(
        os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
        "excel",
    )
    os.makedirs(carpeta_excel, exist_ok=True)
    return os.path.join(carpeta_excel, "turnos.xlsx")
# ==========================================================
# INICIO
# ==========================================================
def inicio(request):
    trabajos = Trabajo.objects.all()[:6]
    return render(
        request,
        "tallermami/inicio.html",
        {"trabajos": trabajos},
    )
# ==========================================================
# LOGIN
# ==========================================================
def login_view(request):
    if request.user.is_authenticated:
        return redirect("inicio")
    if request.method == "POST":
        form = AuthenticationForm(request, data=request.POST)
        if form.is_valid():
            usuario = form.get_user()
            login(request, usuario)
            messages.success(
                request,
                "Iniciaste sesión correctamente.",
            )
            return redirect("inicio")
    else:
        form = AuthenticationForm()
    form.fields["username"].widget.attrs.update(
        {
            "class": "form-control",
            "placeholder": "Tu usuario",
            "autocomplete": "username",
            "autofocus": True,
        }
    )
    form.fields["password"].widget.attrs.update(
        {
            "class": "form-control",
            "placeholder": "Tu contraseña",
            "autocomplete": "current-password",
        }
    )
    return render(
        request,
        "tallermami/login.html",
        {"form": form},
    )
# ==========================================================
# REGISTRO
# ==========================================================
def registro(request):
    if request.user.is_authenticated:
        return redirect("inicio")
    if request.method == "POST":
        form = RegistroForm(request.POST)
        if form.is_valid():
            usuario = form.save()
            PerfilUsuario.objects.get_or_create(
                usuario=usuario,
                defaults={"rol": "cliente"},
            )
            login(request, usuario)
            messages.success(
                request,
                "Tu cuenta fue creada correctamente.",
            )
            return redirect("inicio")
    else:
        form = RegistroForm()
    return render(
        request,
        "tallermami/registro.html",
        {"form": form},
    )
# ==========================================================
# LOGOUT
# ==========================================================
def logout_view(request):
    logout(request)
    messages.success(
        request,
        "Sesión cerrada correctamente.",
    )
    return redirect("inicio")
# ==========================================================
# SOLICITAR TURNO
# ==========================================================
def solicitar_turno(request):
    if request.method == "POST":
        form = SolicitarTurnoForm(request.POST)
        if form.is_valid():
            nombre = form.cleaned_data["nombre"].strip()
            apellido = form.cleaned_data["apellido"].strip()
            email_cliente = form.cleaned_data["email"].strip()
            telefono = form.cleaned_data["telefono"].strip()
            dni = form.cleaned_data["dni"]
            acepta_whatsapp = form.cleaned_data["acepta_whatsapp"]
            # ==============================================
            # CREAR O ACTUALIZAR EL CLIENTE AUTOMÁTICAMENTE
            # ==============================================
            cliente = Cliente.objects.filter(
                email__iexact=email_cliente
            ).first()
            if cliente:
                cliente.nombre = nombre
                cliente.apellido = apellido
                cliente.email = email_cliente
                cliente.telefono = telefono
                cliente.dni = dni
                cliente.acepta_whatsapp = acepta_whatsapp
                cliente.save()
            else:
                cliente = Cliente.objects.create(
                    nombre=nombre,
                    apellido=apellido,
                    dni=dni,
                    email=email_cliente,
                    telefono=telefono,
                    acepta_whatsapp=acepta_whatsapp,
                )
            # ==============================================
            # ASOCIAR EL TURNO A UN USUARIO
            # ==============================================
            if request.user.is_authenticated:
                usuario = request.user
                # Usar el correo ingresado para que
                # las notificaciones lleguen al cliente.
                if usuario.email != email_cliente:
                    usuario.email = email_cliente
                    usuario.save(update_fields=["email"])
            else:
                usuario = User.objects.filter(
                    email__iexact=email_cliente
                ).first()
                if not usuario:
                    base_username = email_cliente.split("@")[0]
                    username = base_username
                    contador = 1
                    while User.objects.filter(
                        username=username
                    ).exists():
                        username = f"{base_username}{contador}"
                        contador += 1
                    usuario = User(
                        username=username,
                        email=email_cliente,
                    )
                    usuario.set_unusable_password()
                    usuario.save()
            if not cliente.usuario_id and not Cliente.objects.filter(
                usuario=usuario
            ).exclude(pk=cliente.pk).exists():
                cliente.usuario = usuario
                cliente.save(update_fields=["usuario"])
            PerfilUsuario.objects.get_or_create(
                usuario=usuario,
                defaults={"rol": "cliente"},
            )
            # ==============================================
            # GUARDAR EL TURNO
            # ==============================================
            turno = form.save(commit=False)
            turno.usuario = usuario
            turno.save()
            vehiculo, _ = Vehiculo.objects.update_or_create(
                patente=turno.patente,
                defaults={
                    "marca": turno.marca,
                    "modelo": turno.modelo,
                    "anio": turno.anio,
                    "cliente": cliente,
                },
            )
            turno.cliente = cliente
            turno.vehiculo_asociado = vehiculo
            turno.save(
                update_fields=[
                    "cliente",
                    "vehiculo_asociado",
                ]
            )
            try:
                guardar_turno_con_disponibilidad(
                    turno,
                    turno.fecha_turno,
                )
            except ValidationError as error:
                turno.delete()
                form.add_error("fecha_turno", error.messages[0])
                return render(
                    request,
                    "tallermami/solicitar_turno.html",
                    {"form": form},
                )
            # ==============================================
            # GUARDAR EN EXCEL
            # ==============================================
            try:
                guardar_turno_en_excel(turno)
            except Exception as error:
                print(f"Error al guardar el turno en Excel: {error}")
            # ==============================================
            # ENVIAR EMAIL DE CONFIRMACIÓN
            # ==============================================
            try:
                enviar_correo_taller(
                    asunto="Solicitud de turno - Taller Molina Henze",
                    mensaje=(
                        f"Hola {nombre},\n\n"
                        "Recibimos correctamente tu solicitud de turno.\n\n"
                        f"Vehículo: {turno.vehiculo}\n"
                        f"Patente: {turno.patente}\n"
                        f"Servicio: {turno.servicio}\n"
                        f"Fecha: {turno.fecha_turno.strftime('%d/%m/%Y %H:%M')}\n\n"
                        "Te esperamos en Taller Molina Henze.\n\n"
                        "Saludos."
                    ),
                    destinatarios=[email_cliente],
                    request=request,
                    fail_silently=False,
                )
            except Exception as error:
                print(f"Error al enviar email: {error}")
            if acepta_whatsapp:
                mensaje = (
                    f"Recibimos tu solicitud de turno para {turno.vehiculo}. "
                    f"Fecha solicitada: "
                    f"{turno.fecha_turno.strftime('%d/%m/%Y %H:%M')}."
                )
                try:
                    resultado_whatsapp = enviar_whatsapp_a_cliente(
                        turno,
                        mensaje,
                    )
                except (
                    ImproperlyConfigured,
                    WhatsAppDeliveryError,
                ):
                    logger.exception(
                        "No se pudo enviar la confirmación de WhatsApp "
                        "del turno %s.",
                        turno.pk,
                    )
                    messages.warning(
                        request,
                        "El turno fue solicitado, pero no se pudo enviar "
                        "la confirmación por WhatsApp. Revisá la "
                        "configuración del servicio.",
                    )
                else:
                    if resultado_whatsapp == "enviado":
                        messages.success(
                            request,
                            "También enviamos la confirmación por WhatsApp.",
                        )
                    elif resultado_whatsapp == "sin_telefono":
                        messages.warning(
                            request,
                            "No se pudo enviar WhatsApp porque no hay un "
                            "teléfono registrado.",
                        )
            return redirect(
                "turno_exitoso",
                turno_id=turno.id,
            )
    else:
        form = SolicitarTurnoForm()
    return render(
        request,
        "tallermami/solicitar_turno.html",
        {"form": form},
    )
# ==========================================================
# TURNO EXITOSO
# ==========================================================
def turno_exitoso(request, turno_id):
    turno = get_object_or_404(Turno, id=turno_id)
    return render(
        request,
        "tallermami/turno_exitoso.html",
        {"turno": turno},
    )
# ==========================================================
# MIS TURNOS
# ==========================================================
@login_required
def mis_turnos(request):
    turnos = (
        Turno.objects.filter(usuario=request.user)
        .exclude(estado="cancelado")
        .select_related("servicio")
    )
    return render(
        request,
        "tallermami/mis_turnos.html",
        {"turnos": turnos},
    )
# ==========================================================
# CANCELAR TURNO
# ==========================================================
@login_required
def cancelar_turno(request, turno_id):
    turno = get_object_or_404(
        Turno,
        id=turno_id,
        usuario=request.user,
    )
    if turno.estado == "cancelado":
        messages.error(
            request,
            "Este turno ya está cancelado.",
        )
        return redirect("mis_turnos")
    fecha_liberada = turno.fecha_turno
    turno.estado = "cancelado"
    campos_actualizados = ["estado"]
    if turno.disponibilidad_id:
        turno.disponibilidad.estado = True
        turno.disponibilidad.save(update_fields=["estado"])
        turno.disponibilidad = None
        campos_actualizados.append("disponibilidad")
    turno.save(update_fields=campos_actualizados)
    # Si el horario liberado ya pasó, no se ofrece adelanto.
    if fecha_liberada <= timezone.now():
        messages.success(
            request,
            "Tu turno fue cancelado correctamente.",
        )
        return redirect("mis_turnos")
    # Buscar turnos posteriores al horario liberado
    candidatos = (
        Turno.objects.filter(
            fecha_turno__gt=fecha_liberada,
            estado="pendiente",
            usuario__isnull=False,
        )
        .exclude(id=turno.id)
        .exclude(usuario__email="")
        .select_related("usuario", "servicio")
        .order_by("fecha_turno")
    )
    # Elegir el primer candidato sin oferta pendiente
    turno_para_ofrecer = None
    for candidato in candidatos:
        tiene_oferta_pendiente = OfertaTurno.objects.filter(
            turno_cliente=candidato,
            estado="pendiente",
        ).exists()
        if not tiene_oferta_pendiente:
            turno_para_ofrecer = candidato
            break
    if turno_para_ofrecer:
        oferta = OfertaTurno.objects.create(
            turno_cliente=turno_para_ofrecer,
            fecha_disponible=fecha_liberada,
        )
        enlace = request.build_absolute_uri(
            reverse(
                "aceptar_oferta_turno",
                args=[oferta.token],
            )
        )
        enviar_correo_taller(
            asunto="¡Podés adelantar tu turno!",
            mensaje=(
                f"Hola {turno_para_ofrecer.usuario.get_full_name() or turno_para_ofrecer.usuario.username},\n\n"
                "Se liberó un horario anterior al de tu turno actual.\n\n"
                "Si querés adelantar tu turno, ingresá al siguiente enlace:\n\n"
                "QUIERO ADELANTAR MI TURNO:\n"
                f"{enlace}\n\n"
                "La disponibilidad está sujeta a confirmación.\n\n"
                "Saludos.\n"
                "Taller Molina Henze"
            ),
            destinatarios=[turno_para_ofrecer.usuario.email],
            request=request,
            fail_silently=False,
        )
    messages.success(
        request,
        "Tu turno fue cancelado correctamente.",
    )
    return redirect("mis_turnos")
# ==========================================================
# ACEPTAR OFERTA DE ADELANTO
# ==========================================================
def aceptar_oferta_turno(request, token):
    oferta = get_object_or_404(OfertaTurno, token=token)
    if oferta.estado != "pendiente":
        return render(
            request,
            "tallermami/oferta_no_disponible.html",
        )
    turno = oferta.turno_cliente
    if turno.estado == "cancelado":
        return render(
            request,
            "tallermami/oferta_no_disponible.html",
        )
    if request.method == "GET":
        return render(
            request,
            "tallermami/oferta_confirmar.html",
            {"oferta": oferta, "turno": turno},
        )
    if request.method != "POST":
        return HttpResponseNotAllowed(["GET", "POST"])
    if request.POST.get("confirmar") != "si":
        return HttpResponse(
            "Debés confirmar que querés adelantar el turno.",
            status=400,
        )

    horario_anterior = turno.fecha_turno
    try:
        with transaction.atomic():
            guardar_turno_con_disponibilidad(
                turno,
                oferta.fecha_disponible,
            )
            oferta.estado = "aceptada"
            oferta.save(update_fields=["estado"])
    except ValidationError:
        return render(
            request,
            "tallermami/oferta_no_disponible.html",
        )

    if turno.usuario and turno.usuario.email:
        enviar_correo_taller(
            asunto="Turno adelantado correctamente",
            mensaje=(
                f"Hola {turno.usuario.get_full_name() or turno.usuario.username},\n\n"
                "Tu turno fue adelantado correctamente.\n\n"
                "Nuevo horario:\n"
                f"{turno.fecha_turno.strftime('%d/%m/%Y %H:%M')}\n\n"
                "Gracias por utilizar nuestro sistema de turnos.\n\n"
                "Taller Molina Henze"
            ),
            destinatarios=[turno.usuario.email],
            request=request,
            fail_silently=True,
        )
    return render(
        request,
        "tallermami/oferta_aceptada.html",
        {
            "turno": turno,
            "horario_anterior": horario_anterior,
        },
    )
# ==========================================================
# ESTADO DE REPARACIÓN
# ==========================================================
def estado_reparacion(request):
    turno = None
    patente_buscada = ""
    if request.method == "POST":
        patente_buscada = (
            request.POST.get("patente", "").strip().upper()
        )
        if patente_buscada:
            turno = (
                Turno.objects.filter(patente__iexact=patente_buscada)
                .select_related("servicio")
                .order_by("-fecha_creacion")
                .first()
            )
            if not turno:
                messages.error(
                    request,
                    "No encontramos ningún vehículo con esa patente.",
                )
    return render(
        request,
        "tallermami/estado_reparacion.html",
        {
            "turno": turno,
            "patente_buscada": patente_buscada,
        },
    )
# ==========================================================
# PDF
# ==========================================================
def descargar_pdf(request, turno_id):
    from .utils import generar_pdf_turno
    turno = get_object_or_404(Turno, id=turno_id)
    return generar_pdf_turno(turno)
# ==========================================================
# EXCEL
# ==========================================================
def exportar_excel(request):
    turnos = Turno.objects.select_related(
        "usuario",
        "servicio",
    ).all()
    return render(
        request,
        "tallermami/exportar_excel.html",
        {"turnos": turnos},
    )
@login_required
def descargar_excel_turnos(request):
    if not (es_jefe_taller(request.user) or es_secretario(request.user)):
        return HttpResponse(
            "No tenés permiso para descargar este archivo.",
            status=403,
        )
    from .utils import crear_excel_turnos
    turnos, busqueda = obtener_turnos_filtrados(request)
    archivo = crear_excel_turnos(turnos)
    response = HttpResponse(
        archivo.getvalue(),
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
    )
    response["Content-Disposition"] = 'attachment; filename="turnos.xlsx"'
    return response
def guardar_turno_en_excel(turno):
    ruta_excel = obtener_ruta_excel()
    if os.path.exists(ruta_excel):
        libro = load_workbook(ruta_excel)
    else:
        libro = Workbook()
        hoja = libro.active
        hoja.title = "Todos los turnos"
    for nombre_hoja in [
        "Todos los turnos",
        "Clientes nuevos",
        "Clientes habituales",
    ]:
        if nombre_hoja not in libro.sheetnames:
            libro.create_sheet(nombre_hoja)
    encabezados = [
        "ID",
        "Cliente",
        "Email",
        "Teléfono",
        "Vehículo",
        "Patente",
        "Servicio",
        "Fecha",
        "Hora",
        "Tipo de cliente",
        "Estado",
        "Avance",
        "Observaciones",
    ]
    for nombre_hoja in libro.sheetnames:
        hoja = libro[nombre_hoja]
        if hoja.max_row == 1 and hoja.cell(1, 1).value is None:
            for columna, encabezado in enumerate(encabezados, start=1):
                celda = hoja.cell(row=1, column=columna)
                celda.value = encabezado
                celda.font = Font(bold=True)
                celda.alignment = Alignment(horizontal="center")
    nombre = ""
    email = ""
    telefono = ""
    if turno.usuario:
        nombre = turno.usuario.get_full_name() or turno.usuario.username
        email = turno.usuario.email
    cantidad_anterior = (
        Turno.objects.filter(usuario=turno.usuario)
        .exclude(id=turno.id)
        .count()
    )
    tipo_cliente = (
        "Cliente habitual" if cantidad_anterior > 0 else "Cliente nuevo"
    )
    fila = [
        turno.id,
        nombre,
        email,
        telefono,
        turno.vehiculo,
        turno.patente,
        str(turno.servicio),
        turno.fecha_turno.strftime("%d/%m/%Y"),
        turno.fecha_turno.strftime("%H:%M"),
        tipo_cliente,
        turno.get_estado_display(),
        f"{turno.porcentaje_reparacion}%",
        turno.observaciones,
    ]
    # Evitar filas duplicadas en "Todos los turnos"
    hoja_todos = libro["Todos los turnos"]
    ids_existentes = [
        hoja_todos.cell(row=fila_excel, column=1).value
        for fila_excel in range(2, hoja_todos.max_row + 1)
    ]
    if turno.id not in ids_existentes:
        hoja_todos.append(fila)
    # Evitar filas duplicadas en la hoja del tipo de cliente
    hoja_cliente = libro[
        "Clientes habituales"
        if tipo_cliente == "Cliente habitual"
        else "Clientes nuevos"
    ]
    ids_existentes_cliente = [
        hoja_cliente.cell(row=fila_excel, column=1).value
        for fila_excel in range(2, hoja_cliente.max_row + 1)
    ]
    if turno.id not in ids_existentes_cliente:
        hoja_cliente.append(fila)
    # Ajustar el ancho de las columnas
    for nombre_hoja in libro.sheetnames:
        hoja = libro[nombre_hoja]
        for columna in range(1, hoja.max_column + 1):
            letra = get_column_letter(columna)
            ancho = 15
            for celda in hoja[letra]:
                if celda.value:
                    ancho = max(ancho, len(str(celda.value)) + 2)
            hoja.column_dimensions[letra].width = min(ancho, 40)
    libro.save(ruta_excel)
def reconstruir_excel():
    ruta_excel = obtener_ruta_excel()
    libro = Workbook()
    hoja_todos = libro.active
    hoja_todos.title = "Todos los turnos"
    hoja_nuevos = libro.create_sheet("Clientes nuevos")
    hoja_habituales = libro.create_sheet("Clientes habituales")
    encabezados = [
        "ID",
        "Cliente",
        "Email",
        "Vehículo",
        "Patente",
        "Servicio",
        "Fecha",
        "Hora",
        "Tipo de cliente",
        "Estado",
        "Avance",
        "Observaciones",
    ]
    for hoja in [hoja_todos, hoja_nuevos, hoja_habituales]:
        for columna, encabezado in enumerate(encabezados, start=1):
            celda = hoja.cell(row=1, column=columna)
            celda.value = encabezado
            celda.font = Font(bold=True)
    turnos = Turno.objects.select_related("usuario", "servicio").all()
    for turno in turnos:
        if turno.usuario:
            cliente = turno.usuario.get_full_name() or turno.usuario.username
            email = turno.usuario.email
        else:
            cliente = "Cliente"
            email = ""
        anteriores = (
            Turno.objects.filter(usuario=turno.usuario)
            .exclude(id=turno.id)
            .count()
        )
        tipo = "Cliente habitual" if anteriores > 0 else "Cliente nuevo"
        fila = [
            turno.id,
            cliente,
            email,
            turno.vehiculo,
            turno.patente,
            str(turno.servicio),
            turno.fecha_turno.strftime("%d/%m/%Y"),
            turno.fecha_turno.strftime("%H:%M"),
            tipo,
            turno.get_estado_display(),
            f"{turno.porcentaje_reparacion}%",
            turno.observaciones,
        ]
        hoja_todos.append(fila)
        if tipo == "Cliente nuevo":
            hoja_nuevos.append(fila)
        else:
            hoja_habituales.append(fila)
    for hoja in [hoja_todos, hoja_nuevos, hoja_habituales]:
        for columna in range(1, hoja.max_column + 1):
            letra = get_column_letter(columna)
            hoja.column_dimensions[letra].width = 20
    libro.save(ruta_excel)
# ==========================================================
# FILTRAR TURNOS PARA LOS PANELES
# ==========================================================
def obtener_turnos_filtrados(request):
    busqueda = request.GET.get("buscar", "").strip()
    turnos = Turno.objects.select_related(
        "usuario",
        "servicio",
        "trabajador",
    ).order_by("-fecha_turno")
    if busqueda:
        emails_clientes = Cliente.objects.filter(
            Q(nombre__icontains=busqueda)
            | Q(apellido__icontains=busqueda)
            | Q(email__icontains=busqueda)
        ).values_list("email", flat=True)
        turnos = turnos.filter(
            Q(patente__icontains=busqueda)
            | Q(vehiculo__icontains=busqueda)
            | Q(usuario__username__icontains=busqueda)
            | Q(usuario__first_name__icontains=busqueda)
            | Q(usuario__last_name__icontains=busqueda)
            | Q(usuario__email__icontains=busqueda)
            | Q(usuario__email__in=emails_clientes)
        ).distinct()
    return turnos, busqueda
# ==========================================================
# PERMISOS DE LOS CARGOS
# ==========================================================
def es_administrador(user):
    return (
        user.is_authenticated
        and (
            user.is_staff
            or PerfilUsuario.objects.filter(
                usuario=user,
                rol="administrador",
            ).exists()
        )
    )


def es_jefe_taller(user):
    return (
        user.is_authenticated
        and (
            es_administrador(user)
            or user.groups.filter(name="Jefe de taller").exists()
        )
    )
def es_secretario(user):
    return (
        user.is_authenticated
        and (
            es_administrador(user)
            or user.groups.filter(name="Secretario").exists()
        )
    )


@login_required(login_url="login")
def panel_jefe_compat(request):
    if not es_jefe_taller(request.user):
        return HttpResponse(
            "No tenés permiso para acceder a esta página.",
            status=403,
        )
    turnos = Turno.objects.select_related(
        "usuario",
        "servicio",
        "trabajador",
    ).order_by("-fecha_turno")
    return render(
        request,
        "tallermami/panel_jefe.html",
        {
            "turnos": turnos,
            "cantidad_turnos": turnos.count(),
            "turnos_pendientes": turnos.filter(
                estado="pendiente"
            ).count(),
        },
    )


@login_required(login_url="login")
def panel_mecanico(request):
    mecanico = Trabajador.objects.filter(usuario=request.user).first()
    es_admin = es_administrador(request.user)
    if mecanico is None and not es_admin:
        messages.error(
            request,
            "Tu cuenta todavía no está asociada a un mecánico. "
            "Pedile al administrador que la vincule en el registro "
            "del trabajador.",
        )
        return redirect("inicio")

    turnos = Turno.objects.exclude(estado="cancelado")
    if not es_admin:
        turnos = turnos.filter(trabajador=mecanico)
    turnos = turnos.select_related(
        "usuario",
        "cliente",
        "servicio",
        "vehiculo_asociado",
    ).order_by("-fecha_turno")
    return render(
        request,
        "tallermami/panel_mecanico.html",
        {
            "mecanico": mecanico,
            "es_admin": es_admin,
            "turnos": turnos,
        },
    )


@login_required(login_url="login")
def mecanico_actualizar_avance(request, turno_id):
    mecanico = Trabajador.objects.filter(usuario=request.user).first()
    es_admin = es_administrador(request.user)
    if mecanico is None and not es_admin:
        messages.error(
            request,
            "Tu cuenta todavía no está asociada a un mecánico. "
            "Pedile al administrador que la vincule en el registro "
            "del trabajador.",
        )
        return redirect("inicio")

    if request.method != "POST":
        return redirect("panel_mecanico")

    turnos = Turno.objects.select_related("usuario", "trabajador")
    if not es_admin:
        turnos = turnos.filter(trabajador=mecanico)
    turno = get_object_or_404(turnos, pk=turno_id)
    form = AvanceReparacionForm(request.POST)
    if not form.is_valid():
        for error in form.errors.get("porcentaje_reparacion", []):
            messages.error(request, error)
        return redirect("panel_mecanico")

    porcentaje_anterior = turno.porcentaje_reparacion
    porcentaje_nuevo = form.cleaned_data["porcentaje_reparacion"]
    if porcentaje_nuevo == porcentaje_anterior:
        messages.info(request, "El porcentaje no tuvo cambios.")
        return redirect("panel_mecanico")

    turno.porcentaje_reparacion = porcentaje_nuevo
    turno.save(update_fields=["porcentaje_reparacion"])
    messages.success(
        request,
        f"Se actualizó el avance del turno #{turno.pk} "
        f"al {porcentaje_nuevo}%.",
    )

    try:
        resultado = notificar_avance_reparacion(
            turno,
            porcentaje_anterior,
            request,
        )
    except Exception:
        logger.exception(
            "No se pudo enviar el aviso por correo del avance "
            "del turno %s desde el panel del mecánico.",
            turno.pk,
        )
        messages.warning(
            request,
            "El avance se guardó, pero no se pudo enviar el correo "
            "al cliente.",
        )
    else:
        if resultado == "sin_correo":
            messages.warning(
                request,
                "El avance se guardó, pero el cliente no tiene "
                "un correo registrado.",
            )
        elif resultado == "sin_enviar":
            messages.warning(
                request,
                "El avance se guardó, pero el servidor no confirmó "
                "el envío del correo.",
            )
        elif resultado == "enviado":
            messages.success(request, "Se notificó al cliente por correo.")

    try:
        resultado_whatsapp = enviar_whatsapp_a_cliente(
            turno,
            mensaje_avance_whatsapp(turno, porcentaje_anterior),
        )
    except (ImproperlyConfigured, WhatsAppDeliveryError):
        logger.exception(
            "No se pudo enviar el aviso por WhatsApp del avance "
            "del turno %s desde el panel del mecánico.",
            turno.pk,
        )
        messages.warning(
            request,
            "El avance se guardó, pero no se pudo enviar el aviso "
            "por WhatsApp.",
        )
    else:
        if resultado_whatsapp == "enviado":
            messages.success(request, "También se notificó al cliente por WhatsApp.")

    return redirect("panel_mecanico")


# ==========================================================
# PANEL DEL JEFE DE TALLER
# ==========================================================
@login_required(login_url="login")
@user_passes_test(es_jefe_taller, login_url="inicio")
def vista_jefe_taller(request):
    trabajos = Trabajo.objects.all().order_by("nombre")
    trabajadores = Trabajador.objects.all().order_by("apellido", "nombre")
    clientes = Cliente.objects.all().order_by("apellido", "nombre")
    turnos, busqueda = obtener_turnos_filtrados(request)
    presupuestos = Presupuesto.objects.all().order_by("-id")
    return render(
        request,
        "tallermami/jefe_taller.html",
        {
            "trabajos": trabajos,
            "trabajadores": trabajadores,
            "clientes": clientes,
            "turnos": turnos,
            "busqueda": busqueda,
            "presupuestos": presupuestos,
        },
    )
@login_required(login_url="login")
@user_passes_test(es_jefe_taller, login_url="inicio")
def vista_jefe_taller(request):
    print(
        "JEFE:",
        request.user.username,
        request.user.is_authenticated,
        es_jefe_taller(request.user),
    )

    trabajos = Trabajo.objects.all().order_by("nombre")
    trabajadores = Trabajador.objects.all().order_by("apellido", "nombre")
    clientes = Cliente.objects.all().order_by("apellido", "nombre")
    turnos, busqueda = obtener_turnos_filtrados(request)

    return render(
        request,
        "tallermami/jefe_taller.html",
        {
            "trabajos": trabajos,
            "trabajadores": trabajadores,
            "clientes": clientes,
            "turnos": turnos,
            "busqueda": busqueda,
        },
    )
# ==========================================================
# PANEL DEL SECRETARIO
# ==========================================================
@login_required(login_url="login")
@user_passes_test(es_secretario, login_url="inicio")
def vista_secretario(request):
    clientes = Cliente.objects.all().order_by("apellido", "nombre")
    turnos, busqueda = obtener_turnos_filtrados(request)
    presupuestos = Presupuesto.objects.all().order_by("-id")
    return render(
        request,
        "tallermami/secretario.html",
        {
            "clientes": clientes,
            "turnos": turnos,
            "busqueda": busqueda,
            "presupuestos": presupuestos,
        },
    )
# ==========================================================
# SECCIONES DEL JEFE DE TALLER
# ==========================================================
@login_required(login_url="login")
@user_passes_test(es_jefe_taller, login_url="inicio")
def jefe_trabajos(request):
    return render(
        request,
        "tallermami/seccion_panel.html",
        {
            "panel": "jefe",
            "seccion": "trabajos",
            "titulo": "Trabajos y servicios",
            "trabajos": Trabajo.objects.all().order_by("nombre"),
        },
    )
@login_required(login_url="login")
@user_passes_test(es_jefe_taller, login_url="inicio")
def jefe_trabajadores(request):
    return render(
        request,
        "tallermami/seccion_panel.html",
        {
            "panel": "jefe",
            "seccion": "trabajadores",
            "titulo": "Trabajadores",
            "trabajadores": Trabajador.objects.all().order_by("apellido", "nombre"),
        },
    )
@login_required(login_url="login")
@user_passes_test(es_jefe_taller, login_url="inicio")
def jefe_clientes(request):
    return render(
        request,
        "tallermami/seccion_panel.html",
        {
            "panel": "jefe",
            "seccion": "clientes",
            "titulo": "Clientes",
            "clientes": Cliente.objects.all().order_by("apellido", "nombre"),
        },
    )
@login_required(login_url="login")
@user_passes_test(es_jefe_taller, login_url="inicio")
def jefe_turnos(request):
    turnos, busqueda = obtener_turnos_filtrados(request)
    return render(
        request,
        "tallermami/seccion_panel.html",
        {
            "panel": "jefe",
            "seccion": "turnos",
            "titulo": "Turnos",
            "turnos": turnos,
            "busqueda": busqueda,
        },
    )


def _modelo_admin(app_label, model_name):
    for model in admin.site._registry:
        options = model._meta
        if options.app_label == app_label and options.model_name == model_name:
            return model, admin.site._registry[model]
    return None, None


def _modelo_contexto(model):
    return {
        "app_label": model._meta.app_label,
        "model_name": model._meta.model_name,
        "verbose_name": str(model._meta.verbose_name),
        "verbose_name_plural": str(model._meta.verbose_name_plural),
        "app_verbose_name": str(model._meta.app_config.verbose_name),
    }


def _columnas_admin(request, model, model_admin):
    columnas = []
    for nombre in model_admin.get_list_display(request):
        if nombre == "__str__":
            etiqueta = str(model._meta.verbose_name)
        else:
            metodo_admin = getattr(model_admin, nombre, None)
            try:
                campo = model._meta.get_field(nombre)
            except FieldDoesNotExist:
                campo = None
            etiqueta = (
                getattr(metodo_admin, "short_description", None)
                or (campo.verbose_name if campo else nombre.replace("_", " "))
            )
        columnas.append({"nombre": nombre, "etiqueta": etiqueta})
    return columnas


def _valor_columna_admin(objeto, model_admin, nombre):
    if nombre == "__str__":
        return str(objeto)
    valor = getattr(model_admin, nombre, None)
    if callable(valor):
        return valor(objeto)
    valor = getattr(objeto, nombre, "")
    return valor() if callable(valor) else valor


def _formulario_jefe_admin(model, model_admin, request, objeto):
    if model is User:
        campos_usuario = (
            "username",
            "first_name",
            "last_name",
            "email",
            "is_active",
            "is_staff",
            "is_superuser",
            "groups",
            "user_permissions",
        )
        if objeto is None:
            class FormularioUsuario(UserCreationForm):
                class Meta(UserCreationForm.Meta):
                    fields = campos_usuario

            return FormularioUsuario

        class FormularioCambioUsuario(UserChangeForm):
            class Meta(UserChangeForm.Meta):
                fields = campos_usuario

        return FormularioCambioUsuario

    campos_solo_lectura = set(
        model_admin.get_readonly_fields(request, objeto)
    )
    campos = []
    for nombre in model_admin.get_fields(request, objeto):
        if nombre in campos_solo_lectura:
            continue
        try:
            campo = model._meta.get_field(nombre)
        except FieldDoesNotExist:
            continue
        if campo.editable and not campo.primary_key:
            campos.append(nombre)

    def formfield_callback(db_field, **kwargs):
        if model is Turno and db_field.name == "fecha_turno":
            return deepcopy(TurnoForm.base_fields["fecha_turno"])
        widget = kwargs.get("widget")
        form_field = model_admin.formfield_for_dbfield(
            db_field,
            request,
            **kwargs,
        )
        if widget is not None:
            form_field.widget = widget
        return form_field

    return modelform_factory(
        model,
        form=model_admin.form,
        fields=campos,
        formfield_callback=formfield_callback,
    )


@login_required(login_url="login")
@user_passes_test(es_jefe_taller, login_url="inicio")
def jefe_admin_inicio(request):
    modelos = sorted(
        admin.site._registry,
        key=lambda model: (
            str(model._meta.app_config.verbose_name),
            str(model._meta.verbose_name_plural),
        ),
    )
    return render(
        request,
        "tallermami/jefe_admin_inicio.html",
        {"modelos": [_modelo_contexto(model) for model in modelos]},
    )


@login_required(login_url="login")
@user_passes_test(es_jefe_taller, login_url="inicio")
def jefe_admin_lista(request, app_label, model_name):
    model, model_admin = _modelo_admin(app_label, model_name)
    if model is None:
        from django.http import Http404
        raise Http404("No existe ese modelo administrable.")

    queryset = model_admin.get_queryset(request)
    acciones = model_admin.get_actions(request)
    if request.method == "POST":
        nombre_accion = request.POST.get("accion", "")
        seleccionados = request.POST.getlist("seleccionado")
        accion = acciones.get(nombre_accion)
        if accion is None:
            messages.error(request, "La acción seleccionada no es válida.")
        elif not seleccionados:
            messages.warning(
                request,
                "Seleccioná al menos un registro para ejecutar la acción.",
            )
        else:
            objetos = queryset.filter(pk__in=seleccionados)
            funcion = accion[0] if isinstance(accion, tuple) else accion.func
            respuesta = funcion(model_admin, request, objetos)
            if respuesta is not None:
                return respuesta
        return redirect(
            "jefe_admin_lista",
            app_label=app_label,
            model_name=model_name,
        )

    busqueda = request.GET.get("q", "").strip()
    search_fields = model_admin.get_search_fields(request)
    if busqueda and search_fields:
        filtros = Q()
        for search_field in search_fields:
            field_name = search_field.lstrip("^=@")
            filtros |= Q(**{f"{field_name}__icontains": busqueda})
        queryset = queryset.filter(filtros).distinct()

    if not queryset.ordered:
        queryset = queryset.order_by(model._meta.pk.name)
    paginator = Paginator(queryset, 50)
    pagina = paginator.get_page(request.GET.get("page"))
    columnas = _columnas_admin(request, model, model_admin)
    filas = [
        {
            "objeto": objeto,
            "valores": [
                _valor_columna_admin(objeto, model_admin, columna["nombre"])
                for columna in columnas
            ],
        }
        for objeto in pagina.object_list
    ]
    return render(
        request,
        "tallermami/jefe_admin_lista.html",
        {
            "modelo": model,
            "modelo_info": _modelo_contexto(model),
            "model_admin": model_admin,
            "acciones": [
                {"nombre": nombre, "etiqueta": etiqueta}
                for nombre, etiqueta in model_admin.get_action_choices(request)
                if nombre
            ],
            "columnas": columnas,
            "filas": filas,
            "pagina": pagina,
            "busqueda": busqueda,
        },
    )


@login_required(login_url="login")
@user_passes_test(es_jefe_taller, login_url="inicio")
def jefe_admin_formulario(request, app_label, model_name, object_id=None):
    model, model_admin = _modelo_admin(app_label, model_name)
    if model is None:
        from django.http import Http404
        raise Http404("No existe ese modelo administrable.")

    objeto = (
        get_object_or_404(model, pk=object_id)
        if object_id is not None
        else None
    )
    cambio = objeto is not None
    form_class = _formulario_jefe_admin(
        model,
        model_admin,
        request,
        objeto,
    )
    form = form_class(
        request.POST or None,
        request.FILES or None,
        instance=objeto,
    )
    if model is Turno and "estado" in form.fields:
        campo_estado = form.fields["estado"]
        campo_estado.choices = Turno.ESTADOS
        campo_estado.widget = forms.Select(
            attrs={
                "class": "form-select jefe-status-select",
                "aria-label": "Estado del turno",
            }
        )
        campo_estado.widget.choices = campo_estado.choices
    if request.method == "POST" and form.is_valid():
        objeto = form.save(commit=False)
        model_admin.save_model(request, objeto, form, cambio)
        form.save_m2m()
        messages.success(
            request,
            f"{model._meta.verbose_name.capitalize()} "
            f"{'actualizado' if cambio else 'creado'} correctamente.",
        )
        return redirect(
            "jefe_admin_lista",
            app_label=app_label,
            model_name=model_name,
        )

    for field in form.fields.values():
        if field.widget.__class__.__name__ == "CheckboxInput":
            field.widget.attrs.setdefault("class", "form-check-input")
        elif field.widget.__class__.__name__ in (
            "Select",
            "SelectMultiple",
            "NullBooleanSelect",
        ):
            field.widget.attrs.setdefault("class", "form-select")
        else:
            field.widget.attrs.setdefault("class", "form-control")

    campos_solo_lectura = []
    if objeto is not None:
        for nombre in model_admin.get_readonly_fields(request, objeto):
            etiqueta = nombre.replace("_", " ").capitalize()
            try:
                etiqueta = model._meta.get_field(nombre).verbose_name
            except FieldDoesNotExist:
                pass
            campos_solo_lectura.append(
                {
                    "etiqueta": etiqueta,
                    "valor": _valor_columna_admin(
                        objeto,
                        model_admin,
                        nombre,
                    ),
                }
            )

    return render(
        request,
        "tallermami/jefe_admin_formulario.html",
        {
            "modelo": model,
            "modelo_info": _modelo_contexto(model),
            "form": form,
            "objeto": objeto,
            "cambio": cambio,
            "campos_solo_lectura": campos_solo_lectura,
        },
    )


@login_required(login_url="login")
@user_passes_test(es_jefe_taller, login_url="inicio")
def jefe_admin_eliminar(request, app_label, model_name, object_id):
    model, _model_admin = _modelo_admin(app_label, model_name)
    if model is None:
        from django.http import Http404
        raise Http404("No existe ese modelo administrable.")
    objeto = get_object_or_404(model, pk=object_id)

    if request.method == "POST":
        if isinstance(objeto, User) and objeto.pk == request.user.pk:
            messages.error(request, "No podés eliminar tu propio usuario.")
            return redirect(
                "jefe_admin_lista",
                app_label=app_label,
                model_name=model_name,
            )
        try:
            objeto.delete()
        except ProtectedError:
            messages.error(
                request,
                "No se puede eliminar este registro porque está "
                "relacionado con otros datos.",
            )
        else:
            messages.success(request, "El registro se eliminó correctamente.")
        return redirect(
            "jefe_admin_lista",
            app_label=app_label,
            model_name=model_name,
        )

    return render(
        request,
        "tallermami/jefe_admin_eliminar.html",
        {
            "modelo": model,
            "modelo_info": _modelo_contexto(model),
            "objeto": objeto,
        },
    )


@login_required(login_url="login")
@user_passes_test(es_jefe_taller, login_url="inicio")
def jefe_admin_contrasena(request, object_id):
    usuario = get_object_or_404(User, pk=object_id)
    form = AdminPasswordChangeForm(
        usuario,
        request.POST or None,
    )
    if request.method == "POST" and form.is_valid():
        form.save()
        messages.success(
            request,
            f"Se actualizó la contraseña de {usuario.username}.",
        )
        return redirect(
            "jefe_admin_lista",
            app_label=User._meta.app_label,
            model_name=User._meta.model_name,
        )
    for field in form.fields.values():
        field.widget.attrs.setdefault("class", "form-control")
    return render(
        request,
        "tallermami/jefe_admin_formulario.html",
        {
            "modelo": User,
            "modelo_info": _modelo_contexto(User),
            "form": form,
            "objeto": usuario,
            "cambio_contrasena": True,
        },
    )


# ==========================================================
# SECCIONES DEL SECRETARIO
# ==========================================================
@login_required(login_url="login")
@user_passes_test(es_secretario, login_url="inicio")
def secretario_clientes(request):
    return render(
        request,
        "tallermami/seccion_panel.html",
        {
            "panel": "secretario",
            "seccion": "clientes",
            "titulo": "Clientes",
            "clientes": Cliente.objects.all().order_by("apellido", "nombre"),
        },
    )

@login_required(login_url="login")
@user_passes_test(es_secretario, login_url="inicio")
def secretario_presupuesto_nuevo(request):
    if request.method == "POST":
        form = PresupuestoForm(request.POST)

        if form.is_valid():
            presupuesto = form.save()
            messages.success(
                request,
                f"Se creó correctamente el presupuesto #{presupuesto.id}."
            )
            return redirect("secretario_presupuestos")
    else:
        form = PresupuestoForm()

    return render(
        request,
        "tallermami/presupuesto_form.html",
        {
            "form": form,
            "titulo": "Crear presupuesto",
            "panel": "secretario",
        },
    )

@login_required(login_url="login")
@user_passes_test(es_secretario, login_url="inicio")
def secretario_turnos(request):
    turnos, busqueda = obtener_turnos_filtrados(request)
    cliente_dni = Cliente.objects.filter(
        email__iexact=OuterRef("usuario__email"),
    ).values("dni")[:1]
    turnos = turnos.annotate(cliente_dni=Subquery(cliente_dni))
    return render(
        request,
        "tallermami/seccion_panel.html",
        {
            "panel": "secretario",
            "seccion": "turnos",
            "titulo": "Turnos",
            "turnos": turnos,
            "busqueda": busqueda,
        },
    )
@login_required(login_url="login")
@user_passes_test(es_secretario, login_url="inicio")

@login_required(login_url="login")
@user_passes_test(es_secretario, login_url="inicio")
def secretario_presupuesto_nuevo(request):
    from django.core.mail import send_mail
    from django.db import transaction
    from django.urls import reverse
    from django.conf import settings

    form = PresupuestoForm(request.POST or None)
    formset = RepuestosPresupuestadosFormSet(
        request.POST or None
    )

    if request.method == "POST":
        if form.is_valid() and formset.is_valid():
            presupuesto = form.save(commit=False)

            try:
                with transaction.atomic():
                    presupuesto.save()

                    formset.instance = presupuesto
                    formset.save()

            except Exception:
                messages.error(
                    request,
                    "No se pudo guardar el presupuesto. "
                    "Revisá los datos e intentá nuevamente."
                )
            else:
                turno = presupuesto.turno
                usuario = turno.usuario
                destinatario = (
                    usuario.email.strip()
                    if usuario and usuario.email
                    else ""
                )

                if not destinatario:
                    messages.warning(
                        request,
                        f"Se guardó el presupuesto #{presupuesto.id}, "
                        "pero el turno no tiene un correo electrónico "
                        "de usuario registrado. No se envió el email."
                    )
                    return redirect("secretario_presupuestos")

                url_aceptar = request.build_absolute_uri(
                    reverse(
                        "responder_presupuesto",
                        kwargs={
                            "token": presupuesto.token,
                            "decision": "aceptar",
                        },
                    )
                )

                url_rechazar = request.build_absolute_uri(
                    reverse(
                        "responder_presupuesto",
                        kwargs={
                            "token": presupuesto.token,
                            "decision": "rechazar",
                        },
                    )
                )

                repuestos = presupuesto.repuestos_estimados.all()

                detalle_repuestos = "\n".join(
                    (
                        f"- {item.nombre}: {item.cantidad} x "
                        f"${item.precio_unitario:.2f} = "
                        f"${item.subtotal():.2f}"
                    )
                    for item in repuestos
                ) or "No se agregaron repuestos."

                mensaje = (
                    f"Hola, te enviamos el presupuesto "
                    f"correspondiente a tu vehículo "
                    f"{turno.vehiculo} (patente {turno.patente}).\n\n"
                    f"Servicio: {turno.servicio}\n"
                    f"Mano de obra: "
                    f"${presupuesto.mano_obra_estimada:.2f}\n\n"
                    f"Repuestos:\n{detalle_repuestos}\n\n"
                    f"Total estimado: "
                    f"${presupuesto.total_estimado():.2f}\n\n"
                    f"Observaciones: "
                    f"{presupuesto.observaciones or 'Sin observaciones'}\n\n"
                    f"Para aceptar el presupuesto, ingresá aquí:\n"
                    f"{url_aceptar}\n\n"
                    f"Para rechazarlo, ingresá aquí:\n"
                    f"{url_rechazar}\n\n"
                    "Taller Molina Henze"
                )

                try:
                    enviado = enviar_correo_taller(
                        asunto=(
                            f"Presupuesto #{presupuesto.id} "
                            "- Taller Molina Henze"
                        ),
                        mensaje=mensaje,
                        destinatarios=[destinatario],
                        request=request,
                        fail_silently=False,
                    )

                    if enviado:
                        presupuesto.fecha_envio = timezone.now()
                        presupuesto.save(
                            update_fields=["fecha_envio"]
                        )
                        messages.success(
                            request,
                            f"Presupuesto #{presupuesto.id} creado "
                            "y enviado correctamente al cliente."
                        )
                    else:
                        messages.warning(
                            request,
                            f"Se guardó el presupuesto #{presupuesto.id}, "
                            "pero Django no confirmó el envío del correo."
                        )

                except Exception:
                    messages.warning(
                        request,
                        f"Se guardó el presupuesto #{presupuesto.id}, "
                        "pero ocurrió un error al enviar el correo. "
                        "Verificá la configuración del email."
                    )

                return redirect("secretario_presupuestos")

    return render(
        request,
        "tallermami/presupuesto_form.html",
        {
            "form": form,
            "formset": formset,
            "titulo": "Crear presupuesto",
            "panel": "secretario",
        },
    )

# ==========================================================
# RESPONDER PRESUPUESTO POR CORREO
# ==========================================================
def responder_presupuesto(request, token, decision):
    if decision not in ("aceptar", "rechazar"):
        return HttpResponse(
            "La decisión solicitada no es válida.",
            status=400,
        )
    presupuesto = get_object_or_404(
        Presupuesto.objects.select_related("turno", "turno__usuario"),
        token=token,
    )
    if presupuesto.estado != "pendiente":
        return render(
            request,
            "tallermami/presupuesto_responder.html",
            {
                "presupuesto": presupuesto,
                "decision": decision,
                "puede_responder": False,
                "resultado": (
                    "Este presupuesto ya fue respondido. "
                    "No se puede modificar la decisión."
                ),
            },
        )
    # GET: mostrar la confirmación
    if request.method == "GET":
        return render(
            request,
            "tallermami/presupuesto_responder.html",
            {
                "presupuesto": presupuesto,
                "decision": decision,
                "puede_responder": True,
                "resultado": "",
            },
        )
    # POST: registrar la decisión
    if request.POST.get("confirmar") != "si":
        return HttpResponse(
            "Debés confirmar tu decisión.",
            status=400,
        )
    if decision == "aceptar":
        presupuesto.estado = "aceptado"
        presupuesto.turno.estado = "presupuesto_aceptado"
        resultado = "Aceptaste el presupuesto correctamente."
    else:
        presupuesto.estado = "rechazado"
        presupuesto.turno.estado = "presupuesto_rechazado"
        resultado = "Rechazaste el presupuesto."
    presupuesto.fecha_respuesta = timezone.now()
    presupuesto.save(update_fields=["estado", "fecha_respuesta"])
    presupuesto.turno.save(update_fields=["estado"])
    return render(
        request,
        "tallermami/presupuesto_responder.html",
        {
            "presupuesto": presupuesto,
            "decision": decision,
            "puede_responder": False,
            "resultado": resultado,
        },
    )
@login_required(login_url="login")
@user_passes_test(es_secretario, login_url="inicio")
def secretario_cliente_nuevo(request):
    form = ClienteForm(request.POST or None)

    if request.method == "POST" and form.is_valid():
        form.save()
        messages.success(request, "Cliente creado correctamente.")
        return redirect("secretario_clientes")

    return render(
        request,
        "tallermami/formulario_gestion.html",
        {
            "titulo": "Crear cliente",
            "form": form,
            "volver_url": "secretario_clientes",
        },
    )

@login_required(login_url="login")
@user_passes_test(es_secretario, login_url="inicio")
def secretario_cliente_editar(request, cliente_id):
    cliente = get_object_or_404(Cliente, pk=cliente_id)

    form = ClienteForm(
        request.POST or None,
        instance=cliente
    )

    if request.method == "POST" and form.is_valid():
        form.save()
        messages.success(
            request,
            "Cliente actualizado correctamente."
        )
        return redirect("secretario_clientes")

    return render(
        request,
        "tallermami/formulario_gestion.html",
        {
            "titulo": "Editar cliente",
            "form": form,
            "volver_url": "secretario_clientes",
        },
    )

@login_required(login_url="login")
@user_passes_test(es_secretario, login_url="inicio")
def secretario_turno_editar(request, turno_id):
    turno = get_object_or_404(Turno, pk=turno_id)
    porcentaje_anterior = turno.porcentaje_reparacion

    form = TurnoForm(
        request.POST or None,
        instance=turno
    )
    if request.method == "POST" and form.is_valid():
        try:
            with transaction.atomic():
                turno_editado = form.save(commit=False)
                guardar_turno_con_disponibilidad(
                    turno_editado,
                    form.cleaned_data["fecha_turno"],
                )
                form.save()
        except ValidationError as error:
            form.add_error("fecha_turno", error.messages[0])
            return render(
                request,
                "tallermami/formulario_gestion.html",
                {
                    "titulo": f"Editar turno #{turno.id}",
                    "form": form,
                    "volver_url": "secretario_turnos",
                },
            )
        messages.success(
            request,
            "Turno actualizado correctamente."
        )
        if porcentaje_anterior != turno.porcentaje_reparacion:
            try:
                resultado = notificar_avance_reparacion(
                    turno,
                    porcentaje_anterior,
                    request,
                )
            except Exception:
                logger.exception(
                    "No se pudo enviar la actualización de avance "
                    "del turno %s.",
                    turno.pk,
                )
                messages.warning(
                    request,
                    "El turno se actualizó, pero no se pudo enviar "
                    "el correo de avance.",
                )
            else:
                if resultado == "sin_correo":
                    messages.warning(
                        request,
                        "El turno se actualizó, pero el cliente no tiene "
                        "un correo electrónico registrado.",
                    )
                elif resultado == "sin_enviar":
                    messages.warning(
                        request,
                        "El turno se actualizó, pero el servidor no "
                        "confirmó el envío del correo de avance.",
                    )
                elif resultado == "enviado":
                    messages.success(
                        request,
                        "Se envió al cliente el nuevo porcentaje "
                        "de reparación.",
                    )
            try:
                resultado_whatsapp = enviar_whatsapp_a_cliente(
                    turno,
                    mensaje_avance_whatsapp(
                        turno,
                        porcentaje_anterior,
                    ),
                )
            except (ImproperlyConfigured, WhatsAppDeliveryError):
                logger.exception(
                    "No se pudo enviar la actualización de WhatsApp "
                    "del turno %s.",
                    turno.pk,
                )
                messages.warning(
                    request,
                    "El turno se actualizó, pero no se pudo enviar "
                    "la actualización por WhatsApp.",
                )
            else:
                if resultado_whatsapp == "enviado":
                    messages.success(
                        request,
                        "También se envió la actualización por WhatsApp.",
                    )
        return redirect("secretario_turnos")

    return render(
        request,
        "tallermami/formulario_gestion.html",
        {
            "titulo": f"Editar turno #{turno.id}",
            "form": form,
            "volver_url": "secretario_turnos",
        },
    )


@login_required(login_url="login")
@user_passes_test(es_secretario, login_url="inicio")
def secretario_turno_whatsapp(request, turno_id):
    turno = get_object_or_404(Turno, pk=turno_id)
    cliente = Cliente.objects.filter(
        email__iexact=turno.usuario.email.strip(),
    ).first() if turno.usuario and turno.usuario.email else None

    texto_inicial = (
        f"Hola {cliente.nombre}, te contactamos desde Taller Molina Henze "
        f"por tu turno #{turno.pk} ({turno.vehiculo})."
        if cliente
        else f"Te contactamos desde Taller Molina Henze por tu turno #{turno.pk}."
    )
    form = EnviarWhatsAppForm(
        request.POST or None,
        initial={"mensaje": texto_inicial},
    )

    if request.method == "POST" and form.is_valid():
        if not cliente or not cliente.acepta_whatsapp:
            messages.error(
                request,
                "El cliente no autorizó recibir mensajes por WhatsApp.",
            )
        elif not cliente.telefono.strip():
            messages.error(
                request,
                "El cliente no tiene un teléfono registrado.",
            )
        else:
            try:
                enviar_whatsapp_a_cliente(
                    turno,
                    form.cleaned_data["mensaje"],
                )
            except (ImproperlyConfigured, WhatsAppDeliveryError):
                logger.exception(
                    "No se pudo enviar un WhatsApp manual para el turno %s.",
                    turno.pk,
                )
                messages.error(
                    request,
                    "No se pudo enviar el WhatsApp. Verificá la conexión, "
                    "las credenciales y la plantilla aprobada de Meta.",
                )
            else:
                messages.success(
                    request,
                    "El mensaje fue aceptado por WhatsApp para su envío.",
                )
                return redirect("secretario_turnos")

    return render(
        request,
        "tallermami/enviar_whatsapp.html",
        {
            "turno": turno,
            "cliente": cliente,
            "form": form,
        },
    )

@login_required(login_url="login")
@user_passes_test(es_secretario, login_url="inicio")
def secretario_presupuestos(request):
    cliente_dni = Cliente.objects.filter(
        email__iexact=OuterRef("turno__usuario__email"),
    ).values("dni")[:1]
    presupuestos = (
        Presupuesto.objects
        .select_related("turno", "turno__usuario")
        .prefetch_related("repuestos_estimados")
        .annotate(cliente_dni=Subquery(cliente_dni))
        .order_by("-id")
    )

    return render(
        request,
        "tallermami/seccion_panel.html",
        {
            "panel": "secretario",
            "seccion": "presupuestos",
            "titulo": "Presupuestos",
            "presupuestos": presupuestos,
        },
    )

@login_required(login_url="login")
@user_passes_test(es_secretario, login_url="inicio")
def secretario_presupuesto_final(request, presupuesto_id):
    from django.conf import settings
    from django.core.mail import send_mail
    from django.db import transaction
    from django.shortcuts import get_object_or_404
    from django.utils import timezone

    presupuesto = get_object_or_404(
        Presupuesto.objects.select_related("turno", "turno__usuario"),
        id=presupuesto_id,
    )

    if presupuesto.estado != "aceptado":
        messages.error(
            request,
            "Solo se puede cargar el presupuesto final de un presupuesto aceptado."
        )
        return redirect("secretario_presupuestos")

    if presupuesto.final_enviado:
        messages.info(
            request,
            "El presupuesto final ya fue enviado al cliente."
        )
        return redirect("secretario_presupuestos")

    form = PresupuestoFinalForm(
        request.POST or None,
        instance=presupuesto,
    )
    formset = RepuestosUtilizadosFormSet(
        request.POST or None,
        instance=presupuesto,
    )

    if request.method == "POST":
        if form.is_valid() and formset.is_valid():
            try:
                with transaction.atomic():
                    presupuesto = form.save(commit=False)
                    presupuesto.fecha_finalizacion = timezone.now()
                    presupuesto.save()

                    formset.instance = presupuesto
                    formset.save()

            except Exception:
                messages.error(
                    request,
                    "No se pudo guardar el presupuesto final. "
                    "Revisá los datos e intentá nuevamente."
                )
                return render(
                    request,
                    "tallermami/presupuesto_final_form.html",
                    {
                        "form": form,
                        "formset": formset,
                        "presupuesto": presupuesto,
                    },
                )

            turno = presupuesto.turno
            usuario = turno.usuario
            destinatario = (
                usuario.email.strip()
                if usuario and usuario.email
                else ""
            )

            repuestos = presupuesto.repuestos_utilizados.all()

            detalle_repuestos = "\n".join(
                f"- {item.nombre}: {item.cantidad} x "
                f"${item.precio_unitario:.2f} = ${item.subtotal():.2f}"
                for item in repuestos
            ) or "No se registraron repuestos utilizados."

            total = presupuesto.total_final()

            if not destinatario:
                messages.warning(
                    request,
                    "Se guardó el presupuesto final, pero el cliente "
                    "no tiene un correo electrónico registrado. "
                    "No se envió el email."
                )
                return redirect("secretario_presupuestos")

            mensaje = (
                f"Hola {usuario.get_full_name() or usuario.username}.\n\n"
                "Te informamos que finalizamos el trabajo de tu vehículo "
                "en Taller Molina Henze.\n\n"
                f"Vehículo: {turno.vehiculo}\n"
                f"Patente: {turno.patente}\n"
                f"Servicio: {turno.servicio}\n\n"
                "REPUESTOS UTILIZADOS:\n"
                f"{detalle_repuestos}\n\n"
                f"Mano de obra final: ${presupuesto.mano_obra_final:.2f}\n"
                f"TOTAL FINAL: ${total:.2f}\n\n"
                "Detalle del trabajo realizado:\n"
                f"{presupuesto.detalle_final or 'Sin observaciones adicionales.'}\n\n"
                f"Fecha de finalización: "
                f"{presupuesto.fecha_finalizacion:%d/%m/%Y %H:%M}\n\n"
                "Gracias por confiar en Taller Molina Henze."
            )

            try:
                enviado = enviar_correo_taller(
                    asunto=(
                        f"Trabajo finalizado y detalle de pago "
                        f"- Presupuesto #{presupuesto.id}"
                    ),
                    mensaje=mensaje,
                    destinatarios=[destinatario],
                    request=request,
                    fail_silently=False,
                )

                if enviado:
                    presupuesto.final_enviado = True
                    presupuesto.save(update_fields=["final_enviado"])
                    messages.success(
                        request,
                        "Presupuesto final guardado y correo enviado correctamente."
                    )
                else:
                    messages.warning(
                        request,
                        "El presupuesto se guardó, pero el correo no pudo enviarse."
                    )

            except Exception:
                messages.warning(
                    request,
                    "El presupuesto se guardó, pero ocurrió un error "
                    "al enviar el correo."
                )

            return redirect("secretario_presupuestos")

    return render(
        request,
        "tallermami/presupuesto_final_form.html",
        {
            "form": form,
            "formset": formset,
            "presupuesto": presupuesto,
        },
    )
