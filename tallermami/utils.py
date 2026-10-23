
import json
from io import BytesIO
from urllib.parse import urlencode
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen
from xml.sax.saxutils import escape as xml_escape

from django.conf import settings
from django.core.mail import EmailMultiAlternatives
from django.core.exceptions import ImproperlyConfigured
from django.http import HttpResponse
from django.db.models import Count
from django.template.defaultfilters import linebreaksbr
from django.urls import reverse
from django.utils.html import conditional_escape, urlize

from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

from .models import Turno, Cliente


class WhatsAppDeliveryError(Exception):
    pass


def obtener_cliente_turno(turno):
    if not turno.usuario or not turno.usuario.email:
        return None
    return Cliente.objects.filter(
        email__iexact=turno.usuario.email.strip(),
    ).first()


def normalizar_telefono_whatsapp(telefono):
    telefono = telefono.strip()
    solo_digitos = "".join(caracter for caracter in telefono if caracter.isdigit())

    if telefono.startswith("+"):
        if not 8 <= len(solo_digitos) <= 15:
            raise WhatsAppDeliveryError(
                "El teléfono de WhatsApp debe tener entre 8 y 15 dígitos."
            )
        return solo_digitos

    if solo_digitos.startswith("00"):
        solo_digitos = solo_digitos[2:]
    if solo_digitos.startswith("0"):
        solo_digitos = solo_digitos[1:]
    if solo_digitos.startswith("549") and len(solo_digitos) in range(12, 14):
        return solo_digitos
    if solo_digitos.startswith("54") and len(solo_digitos) in range(11, 13):
        return f"549{solo_digitos[2:]}"
    if len(solo_digitos) == 10:
        return f"549{solo_digitos}"

    raise WhatsAppDeliveryError(
        "Ingresá el teléfono en formato internacional, por ejemplo "
        "+54 9 351 123 4567."
    )


def enviar_whatsapp_taller(telefono, mensaje):
    access_token = getattr(settings, "WHATSAPP_ACCESS_TOKEN", "")
    phone_number_id = getattr(settings, "WHATSAPP_PHONE_NUMBER_ID", "")
    if not access_token or not phone_number_id:
        raise ImproperlyConfigured(
            "WhatsApp no está configurado. Definí "
            "WHATSAPP_ACCESS_TOKEN y WHATSAPP_PHONE_NUMBER_ID."
        )

    numero_destino = normalizar_telefono_whatsapp(telefono)
    version = getattr(settings, "WHATSAPP_GRAPH_API_VERSION", "v22.0")
    template_name = getattr(
        settings,
        "WHATSAPP_TEMPLATE_NAME",
        "taller_mensaje",
    )
    language = getattr(
        settings,
        "WHATSAPP_TEMPLATE_LANGUAGE",
        "es_AR",
    )
    url_pagina = getattr(
        settings,
        "PUBLIC_SITE_URL",
        "http://127.0.0.1:8000",
    ).rstrip("/") + "/"

    payload = {
        "messaging_product": "whatsapp",
        "recipient_type": "individual",
        "to": numero_destino,
        "type": "template",
        "template": {
            "name": template_name,
            "language": {"code": language},
            "components": [
                {
                    "type": "body",
                    "parameters": [
                        {"type": "text", "text": mensaje},
                        {"type": "text", "text": url_pagina},
                    ],
                }
            ],
        },
    }
    request = Request(
        (
            f"https://graph.facebook.com/{version}/"
            f"{phone_number_id}/messages"
        ),
        data=json.dumps(payload).encode("utf-8"),
        headers={
            "Authorization": f"Bearer {access_token}",
            "Content-Type": "application/json",
        },
        method="POST",
    )
    try:
        with urlopen(request, timeout=15) as response:
            resultado = json.loads(response.read().decode("utf-8"))
    except HTTPError as error:
        detalle = error.read().decode("utf-8", errors="replace")
        raise WhatsAppDeliveryError(
            f"WhatsApp rechazó el mensaje (HTTP {error.code}): {detalle}"
        ) from error
    except (URLError, TimeoutError) as error:
        raise WhatsAppDeliveryError(
            f"No se pudo conectar con WhatsApp: {error}"
        ) from error
    except (UnicodeDecodeError, json.JSONDecodeError) as error:
        raise WhatsAppDeliveryError(
            "WhatsApp devolvió una respuesta que no se pudo interpretar."
        ) from error

    if not resultado.get("messages"):
        raise WhatsAppDeliveryError(
            "WhatsApp no confirmó la aceptación del mensaje."
        )
    return resultado


def enviar_whatsapp_a_cliente(turno, mensaje):
    cliente = obtener_cliente_turno(turno)
    if not cliente or not cliente.acepta_whatsapp:
        return "sin_consentimiento"
    if not cliente.telefono.strip():
        return "sin_telefono"
    enviar_whatsapp_taller(cliente.telefono, mensaje)
    return "enviado"


def enviar_correo_taller(
    asunto,
    mensaje,
    destinatarios,
    request=None,
    fail_silently=False,
):
    if request:
        url_pagina = request.build_absolute_uri("/")
    else:
        url_pagina = getattr(
            settings,
            "PUBLIC_SITE_URL",
            "http://127.0.0.1:8000",
        ).rstrip("/") + "/"

    mensaje_texto = (
        f"{mensaje.rstrip()}\n\n"
        f"Conocé más en Taller Molina Henze: {url_pagina}"
    )
    contenido_html = linebreaksbr(urlize(conditional_escape(mensaje)))
    mensaje_html = f"""
        <!doctype html>
        <html lang="es">
          <body style="margin:0;padding:0;background-color:#eef2f7;
                       font-family:Arial,Helvetica,sans-serif;color:#243247;">
            <div style="display:none;max-height:0;overflow:hidden;opacity:0;">
              Información importante de Taller Molina Henze.
            </div>
            <table role="presentation" width="100%" cellspacing="0"
                   cellpadding="0" border="0"
                   style="background-color:#eef2f7;padding:36px 12px;">
              <tr>
                <td align="center">
                  <table role="presentation" width="620" cellspacing="0"
                         cellpadding="0" border="0"
                         style="width:100%;max-width:620px;background:#ffffff;
                                border:1px solid #dce4ee;">
                    <tr>
                      <td style="padding:25px 32px;background:#0d2b4d;
                                 border-bottom:4px solid #0d6efd;">
                        <div style="font-size:10px;line-height:15px;
                                    letter-spacing:2px;text-transform:uppercase;
                                    color:#b9cbe0;">
                          Servicio integral automotriz
                        </div>
                        <div style="margin-top:5px;color:#ffffff;
                                    font-size:23px;line-height:30px;
                                    font-weight:bold;">
                          Taller Molina Henze
                        </div>
                      </td>
                    </tr>
                    <tr>
                      <td style="padding:34px 32px 30px;
                                 font-size:15px;line-height:1.75;
                                 color:#334155;">
                        {contenido_html}
                        <table role="presentation" cellspacing="0"
                               cellpadding="0" border="0"
                               style="margin-top:28px;">
                          <tr>
                            <td bgcolor="#0d6efd"
                                style="border-radius:4px;">
                              <a href="{conditional_escape(url_pagina)}"
                                 style="display:inline-block;padding:13px 22px;
                                        color:#ffffff;text-decoration:none;
                                        font-size:14px;font-weight:bold;">
                                Visitar nuestro sitio
                              </a>
                            </td>
                          </tr>
                        </table>
                      </td>
                    </tr>
                    <tr>
                      <td style="padding:20px 32px;background:#f7f9fc;
                                 border-top:1px solid #e3e9f0;">
                        <div style="font-size:12px;line-height:19px;
                                    color:#526174;font-weight:bold;">
                          Taller Molina Henze
                        </div>
                        <div style="font-size:11px;line-height:18px;
                                    color:#718096;">
                          Servicio integral de mecánica automotriz
                        </div>
                        <a href="{conditional_escape(url_pagina)}"
                           style="font-size:11px;line-height:20px;
                                  color:#0d6efd;text-decoration:none;">
                          {conditional_escape(url_pagina)}
                        </a>
                      </td>
                    </tr>
                  </table>
                  <div style="padding:16px 8px 0;color:#8491a3;
                              font-size:10px;line-height:16px;">
                    Este correo fue enviado por Taller Molina Henze.
                  </div>
                </td>
              </tr>
            </table>
          </body>
        </html>
    """
    email = EmailMultiAlternatives(
        subject=asunto,
        body=mensaje_texto,
        from_email=settings.DEFAULT_FROM_EMAIL,
        to=destinatarios,
    )
    email.attach_alternative(mensaje_html, "text/html")
    return email.send(fail_silently=fail_silently)


def notificar_avance_reparacion(
    turno,
    porcentaje_anterior,
    request=None,
):
    if porcentaje_anterior == turno.porcentaje_reparacion:
        return "sin_cambios"

    email = (
        turno.usuario.email.strip()
        if turno.usuario and turno.usuario.email
        else ""
    )
    if not email:
        return "sin_correo"

    url_estado = reverse("estado_reparacion")
    url_estado = f"{url_estado}?{urlencode({'patente': turno.patente})}"
    if request:
        url_estado = request.build_absolute_uri(url_estado)

    enviado = enviar_correo_taller(
        asunto=f"Actualización de reparación - Turno #{turno.id}",
        mensaje=(
            f"Hola {turno.usuario.get_full_name() or turno.usuario.username},\n\n"
            f"El avance de reparación de tu vehículo "
            f"{turno.vehiculo} (patente {turno.patente}) "
            f"se actualizó del {porcentaje_anterior}% "
            f"al {turno.porcentaje_reparacion}%.\n\n"
            "Podés consultar el estado actualizado ingresando aquí:\n"
            f"{url_estado}\n\n"
            "Taller Molina Henze"
        ),
        destinatarios=[email],
        request=request,
        fail_silently=False,
    )
    return "enviado" if enviado else "sin_enviar"


def mensaje_avance_whatsapp(turno, porcentaje_anterior):
    return (
        f"Actualización de reparación del turno #{turno.pk}: "
        f"{turno.vehiculo} ({turno.patente}) pasó de "
        f"{porcentaje_anterior}% a {turno.porcentaje_reparacion}%."
    )


# ==========================================================
# GENERAR COMPROBANTE PDF
# ==========================================================

def generar_pdf_turno(turno):
    buffer = BytesIO()
    color_marca = colors.HexColor("#0D2B4D")
    color_acento = colors.HexColor("#0D6EFD")
    color_texto = colors.HexColor("#172B42")
    color_suave = colors.HexColor("#F2F6FB")

    documento = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=2 * cm,
        leftMargin=2 * cm,
        topMargin=3.4 * cm,
        bottomMargin=2.5 * cm,
    )

    estilos = getSampleStyleSheet()

    titulo = ParagraphStyle(
        "TituloTaller",
        parent=estilos["Title"],
        alignment=TA_LEFT,
        fontName="Helvetica-Bold",
        fontSize=19,
        leading=23,
        spaceAfter=5,
        textColor=color_marca,
    )

    subtitulo = ParagraphStyle(
        "SubtituloTaller",
        parent=estilos["BodyText"],
        fontSize=9,
        leading=13,
        spaceAfter=14,
        textColor=colors.HexColor("#64748B"),
    )

    encabezado_seccion = ParagraphStyle(
        "EncabezadoSeccionTaller",
        parent=estilos["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=11,
        leading=14,
        spaceBefore=10,
        spaceAfter=8,
        textColor=color_marca,
    )

    valor = ParagraphStyle(
        "ValorTaller",
        parent=estilos["BodyText"],
        fontSize=9,
        leading=13,
        textColor=color_texto,
        wordWrap="CJK",
    )

    etiqueta = ParagraphStyle(
        "EtiquetaTaller",
        parent=valor,
        fontName="Helvetica-Bold",
        fontSize=8,
        leading=11,
        textColor=colors.HexColor("#526174"),
    )

    def dibujar_membrete(canvas, doc):
        canvas.saveState()
        ancho, alto = A4
        canvas.setFillColor(color_marca)
        canvas.rect(0, alto - 1.9 * cm, ancho, 1.9 * cm, fill=1, stroke=0)
        canvas.setFillColor(colors.white)
        canvas.setFont("Helvetica-Bold", 12)
        canvas.drawString(2 * cm, alto - 0.95 * cm, "TALLER MOLINA HENZE")
        canvas.setFont("Helvetica", 8)
        canvas.drawRightString(
            ancho - 2 * cm,
            alto - 0.95 * cm,
            "COMPROBANTE DE TURNO",
        )
        canvas.setStrokeColor(color_acento)
        canvas.setLineWidth(3)
        canvas.line(2 * cm, alto - 2.05 * cm, ancho - 2 * cm, alto - 2.05 * cm)
        canvas.setStrokeColor(colors.HexColor("#DCE4ED"))
        canvas.setLineWidth(0.5)
        canvas.line(2 * cm, 1.8 * cm, ancho - 2 * cm, 1.8 * cm)
        canvas.setFillColor(colors.HexColor("#64748B"))
        canvas.setFont("Helvetica", 8)
        canvas.drawString(
            2 * cm,
            1.25 * cm,
            "Taller Molina Henze - Servicio integral automotriz",
        )
        canvas.drawRightString(ancho - 2 * cm, 1.25 * cm, f"Página {doc.page}")
        canvas.restoreState()

    def parrafo_seguro(contenido, estilo):
        return Paragraph(xml_escape(str(contenido or "—")), estilo)

    contenido = [
        Paragraph("Comprobante de solicitud de turno", titulo),
        Paragraph(
            "Guardá este documento como constancia de los datos de tu solicitud.",
            subtitulo,
        ),
    ]

    usuario = turno.usuario

    if usuario:
        cliente_nombre = (
            usuario.get_full_name()
            or usuario.username
            or "Cliente"
        )
        email = usuario.email or ""
    else:
        cliente_nombre = "Cliente"
        email = "No informado"

    fecha = turno.fecha_turno
    vehiculo = " ".join(
        dato for dato in (turno.marca, turno.modelo) if dato
    ) or turno.vehiculo or "No informado"
    if turno.anio:
        vehiculo = f"{vehiculo} ({turno.anio})"

    campos = [
        ("Número de turno", f"#{turno.id}"),
        ("Cliente", cliente_nombre),
        ("Correo electrónico", email),
        ("Vehículo", vehiculo),
        ("Patente", turno.patente or "No informada"),
        ("Servicio", str(turno.servicio or "No informado")),
        ("Fecha", fecha.strftime("%d/%m/%Y")),
        ("Hora", fecha.strftime("%H:%M")),
        ("Estado", turno.get_estado_display()),
    ]
    tabla_datos = [
        [
            parrafo_seguro(nombre, etiqueta),
            parrafo_seguro(valor_campo, valor),
        ]
        for nombre, valor_campo in campos
    ]
    tabla = Table(
        tabla_datos,
        colWidths=[5 * cm, 12 * cm],
        hAlign="LEFT",
    )
    tabla.setStyle(
        TableStyle([
            ("BACKGROUND", (0, 0), (0, -1), color_suave),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 11),
            ("RIGHTPADDING", (0, 0), (-1, -1), 11),
            ("TOPPADDING", (0, 0), (-1, -1), 9),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 9),
            ("LINEBELOW", (0, 0), (-1, -1), 0.5, colors.HexColor("#DCE4ED")),
            ("BOX", (0, 0), (-1, -1), 0.8, colors.HexColor("#DCE4ED")),
        ])
    )

    contenido.extend([
        tabla,
        Spacer(1, 16),
        Paragraph("Avance de reparación", encabezado_seccion),
    ])

    porcentaje = max(0, min(100, int(turno.porcentaje_reparacion or 0)))
    ancho_avance = 17 * cm
    ancho_completado = ancho_avance * porcentaje / 100
    ancho_pendiente = ancho_avance - ancho_completado
    if ancho_completado == 0:
        ancho_completado, ancho_pendiente = 0.01, ancho_avance - 0.01
    elif ancho_pendiente == 0:
        ancho_completado, ancho_pendiente = ancho_avance - 0.01, 0.01

    barra = Table(
        [["", ""]],
        colWidths=[ancho_completado, ancho_pendiente],
        rowHeights=[0.3 * cm],
        hAlign="LEFT",
    )
    barra.setStyle(
        TableStyle([
            ("BACKGROUND", (0, 0), (0, 0), color_acento),
            ("BACKGROUND", (1, 0), (1, 0), colors.HexColor("#E8EEF5")),
            ("LEFTPADDING", (0, 0), (-1, -1), 0),
            ("RIGHTPADDING", (0, 0), (-1, -1), 0),
            ("TOPPADDING", (0, 0), (-1, -1), 0),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
        ])
    )
    contenido.extend([
        barra,
        Paragraph(f"{porcentaje}% completado", valor),
        Spacer(1, 10),
        Paragraph("Observaciones", encabezado_seccion),
        Paragraph(
            xml_escape(turno.observaciones or "Sin observaciones.").replace(
                "\n",
                "<br/>",
            ),
            valor,
        ),
    ])

    documento.build(contenido, onFirstPage=dibujar_membrete, onLaterPages=dibujar_membrete)
    buffer.seek(0)

    response = HttpResponse(
        buffer.getvalue(),
        content_type="application/pdf",
    )

    response["Content-Disposition"] = (
        f'attachment; filename="comprobante_turno_{turno.id}.pdf"'
    )

    return response


# ==========================================================
# CREAR EXCEL DE TURNOS
# ==========================================================

def crear_excel_turnos(turnos=None):
    buffer = BytesIO()
    libro = Workbook()

    hoja_todos = libro.active
    hoja_todos.title = "Todos los turnos"

    hoja_nuevos = libro.create_sheet("Clientes nuevos")
    hoja_habituales = libro.create_sheet("Clientes habituales")

    encabezados = [
        "ID",
        "Cliente",
        "Email",
        "Teléfono",
        "Vehículo",
        "Patente",
        "Servicio",
        "Fecha",
        "Hora",
        "Tipo de cliente",
        "Estado",
        "Avance",
        "Observaciones",
    ]

    hojas = [
        hoja_todos,
        hoja_nuevos,
        hoja_habituales,
    ]

    # Encabezados y formato
    for hoja in hojas:
        hoja.append(encabezados)
        hoja.freeze_panes = "A2"

        for celda in hoja[1]:
            celda.font = Font(
                bold=True,
                color="FFFFFF",
            )
            celda.fill = PatternFill(
                fill_type="solid",
                fgColor="0D2B4D",
            )
            celda.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

    # Si no se proporciona un conjunto de turnos,
    # se exportan todos los registros.
    if turnos is None:
        turnos = Turno.objects.all()

    turnos = turnos.select_related(
        "usuario",
        "servicio",
    ).order_by("-fecha_turno")

    # Cantidad de turnos por usuario.
    cantidades = dict(
        Turno.objects.values("usuario_id")
        .annotate(cantidad=Count("id"))
        .values_list("usuario_id", "cantidad")
    )

    # Buscar teléfonos de los clientes registrados.
    emails = {
        turno.usuario.email.strip().lower()
        for turno in turnos
        if turno.usuario and turno.usuario.email
    }

    clientes_por_email = {
        cliente.email.strip().lower(): cliente
        for cliente in Cliente.objects.filter(
            email__in=emails
        )
        if cliente.email
    }

    # Completar las tres hojas.
    for turno in turnos:
        usuario = turno.usuario

        if usuario:
            nombre = (
                usuario.get_full_name()
                or usuario.username
                or "Cliente"
            )
            email = usuario.email or ""

            cliente_registrado = clientes_por_email.get(
                email.strip().lower()
            )

            telefono = (
                cliente_registrado.telefono
                if cliente_registrado
                else ""
            )

            cantidad = cantidades.get(usuario.id, 0)
        else:
            nombre = "Cliente"
            email = ""
            telefono = ""
            cantidad = 0

        tipo_cliente = (
            "Cliente habitual"
            if cantidad > 1
            else "Cliente nuevo"
        )

        fecha = turno.fecha_turno

        fila = [
            turno.id,
            nombre,
            email,
            telefono or "",
            turno.vehiculo or "",
            turno.patente or "",
            str(turno.servicio),
            fecha.strftime("%d/%m/%Y"),
            fecha.strftime("%H:%M"),
            tipo_cliente,
            turno.get_estado_display(),
            f"{turno.porcentaje_reparacion or 0}%",
            turno.observaciones or "",
        ]

        hoja_todos.append(fila)

        if tipo_cliente == "Cliente habitual":
            hoja_habituales.append(fila)
        else:
            hoja_nuevos.append(fila)

    # Ajustar ancho de columnas y activar filtros.
    for hoja in hojas:
        hoja.auto_filter.ref = hoja.dimensions

        for columna in hoja.columns:
            letra = get_column_letter(
                columna[0].column
            )

            maximo = max(
                (
                    len(str(celda.value))
                    for celda in columna
                    if celda.value is not None
                ),
                default=0,
            )

            hoja.column_dimensions[letra].width = min(
                maximo + 2,
                40,
            )

    libro.save(buffer)
    buffer.seek(0)

    return buffer
