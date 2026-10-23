
import logging

from django.contrib import admin, messages
from django.core.exceptions import ImproperlyConfigured
from django.db import transaction
from django.http import HttpResponse
from django.urls import reverse
from django.utils import timezone

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from .forms import TurnoAdminForm
from .services import guardar_turno_con_disponibilidad
from .utils import (
    WhatsAppDeliveryError,
    enviar_correo_taller,
    enviar_whatsapp_a_cliente,
    mensaje_avance_whatsapp,
    notificar_avance_reparacion,
)
from .models import (
    Cliente,
    Disponibilidad,
    Mecanico,
    OfertaTurno,
    OrdenTrabajo,
    PerfilUsuario,
    Presupuesto,
    Repuesto,
    RepuestoPresupuestado,
    RepuestoUtilizado,
    Trabajo,
    Trabajador,
    TrabajoRealizado,
    Turno,
    Vehiculo,
)

logger = logging.getLogger(__name__)


# ==========================================================
# FUNCIONES AUXILIARES
# ==========================================================

def obtener_email_cliente(turno):
    """Busca el email asociado al turno."""

    if turno.usuario and turno.usuario.email:
        email = turno.usuario.email.strip()
        if email:
            return email

    if turno.usuario:
        nombre = turno.usuario.first_name.strip()
        apellido = turno.usuario.last_name.strip()

        if nombre and apellido:
            clientes = Cliente.objects.filter(
                nombre__iexact=nombre,
                apellido__iexact=apellido,
            ).exclude(email="")

            if clientes.count() == 1:
                return clientes.first().email.strip()

    return ""


# ==========================================================
# EXPORTAR TURNOS A EXCEL
# ==========================================================

@admin.action(description="Exportar turnos seleccionados a Excel")
def exportar_turnos_excel(modeladmin, request, queryset):
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Turnos"

    encabezados = [
        "ID",
        "Cliente",
        "Email",
        "Vehículo",
        "Patente",
        "Servicio",
        "Trabajador",
        "Fecha",
        "Hora",
        "Estado",
        "Avance",
        "Observaciones",
    ]

    hoja.append(encabezados)

    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill(
            fill_type="solid",
            fgColor="0D2B4D",
        )
        celda.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    turnos = queryset.select_related(
        "usuario",
        "servicio",
        "trabajador",
    )

    for turno in turnos:
        usuario = turno.usuario

        if usuario:
            nombre_cliente = (
                usuario.get_full_name()
                or usuario.username
            )
            email = usuario.email or ""
        else:
            nombre_cliente = "Sin usuario"
            email = ""

        hoja.append([
            turno.id,
            nombre_cliente,
            email,
            turno.vehiculo or "",
            turno.patente or "",
            str(turno.servicio),
            str(turno.trabajador) if turno.trabajador else "Sin asignar",
            turno.fecha_turno.strftime("%d/%m/%Y"),
            turno.fecha_turno.strftime("%H:%M"),
            turno.get_estado_display(),
            f"{turno.porcentaje_reparacion or 0}%",
            turno.observaciones or "",
        ])

    for columna in hoja.columns:
        letra = get_column_letter(columna[0].column)

        maximo = max(
            (
                len(str(celda.value))
                for celda in columna
                if celda.value is not None
            ),
            default=0,
        )

        hoja.column_dimensions[letra].width = min(
            maximo + 2,
            40,
        )

    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = hoja.dimensions

    response = HttpResponse(
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )

    response["Content-Disposition"] = (
        'attachment; filename="turnos_seleccionados.xlsx"'
    )

    libro.save(response)
    return response


# ==========================================================
# TRABAJOS
# ==========================================================

@admin.register(Trabajo)
class TrabajoAdmin(admin.ModelAdmin):
    list_display = ("nombre", "precio_estimado")
    search_fields = ("nombre", "descripcion")
    ordering = ("nombre",)


# ==========================================================
# TRABAJADORES
# ==========================================================

@admin.register(Trabajador)
class TrabajadorAdmin(admin.ModelAdmin):
    list_display = (
        "nombre",
        "apellido",
        "usuario",
        "especialidad",
        "telefono",
    )

    list_filter = ("especialidad",)
    search_fields = (
        "nombre",
        "apellido",
        "especialidad",
        "usuario__username",
        "usuario__email",
    )
    ordering = ("apellido", "nombre")
    filter_horizontal = ("trabajos",)
    fields = (
        "usuario",
        "nombre",
        "apellido",
        "especialidad",
        "telefono",
        "trabajos",
    )


@admin.register(Mecanico)
class MecanicoAdmin(TrabajadorAdmin):
    pass


# ==========================================================
# CLIENTES
# ==========================================================

@admin.register(Cliente)
class ClienteAdmin(admin.ModelAdmin):
    list_display = (
        "nombre",
        "apellido",
        "dni",
        "acepta_whatsapp",
        "email",
        "telefono",
    )

    search_fields = (
        "nombre",
        "apellido",
        "dni",
        "email",
        "telefono",
    )

    ordering = ("apellido", "nombre")


# ==========================================================
# TURNOS
# ==========================================================

@admin.register(Turno)
class TurnoAdmin(admin.ModelAdmin):
    form = TurnoAdminForm

    list_display = (
        "id",
        "mostrar_cliente",
        "marca",
        "modelo",
        "anio",
        "vehiculo",
        "patente",
        "servicio",
        "trabajador",
        "fecha_turno",
        "estado",
        "porcentaje_reparacion",
    )

    list_filter = (
        "estado",
        "fecha_turno",
        "servicio",
        "trabajador",
    )

    search_fields = (
        "usuario__username",
        "usuario__first_name",
        "usuario__last_name",
        "usuario__email",
        "vehiculo",
        "marca",
        "modelo",
        "patente",
        "servicio__nombre",
        "trabajador__nombre",
        "trabajador__apellido",
    )

    list_editable = (
        "trabajador",
        "estado",
    )

    ordering = ("-fecha_turno",)

    fields = (
        "usuario",
        "cliente",
        "vehiculo_asociado",
        "disponibilidad",
        "marca",
        "modelo",
        "anio",
        "vehiculo",
        "patente",
        "servicio",
        "trabajador",
        "fecha_turno",
        "estado",
        "porcentaje_reparacion",
        "observaciones",
    )

    readonly_fields = ("vehiculo", "disponibilidad")

    actions = (exportar_turnos_excel,)

    @admin.display(description="Cliente")
    def mostrar_cliente(self, obj):
        if obj.usuario:
            return (
                obj.usuario.get_full_name()
                or obj.usuario.username
            )
        return "Sin usuario"

    def save_model(self, request, obj, form, change):
        porcentaje_anterior = None
        if change:
            porcentaje_anterior = (
                Turno.objects.filter(pk=obj.pk)
                .values_list("porcentaje_reparacion", flat=True)
                .first()
            )
        with transaction.atomic():
            super().save_model(request, obj, form, change)
            guardar_turno_con_disponibilidad(
                obj,
                obj.fecha_turno,
            )

        if (
            porcentaje_anterior is not None
            and porcentaje_anterior != obj.porcentaje_reparacion
        ):
            try:
                resultado = notificar_avance_reparacion(
                    obj,
                    porcentaje_anterior,
                    request,
                )
            except Exception:
                logger.exception(
                    "No se pudo enviar la actualización de avance "
                    "del turno %s desde Django Admin.",
                    obj.pk,
                )
                self.message_user(
                    request,
                    "El turno se guardó, pero no se pudo enviar "
                    "el correo de avance.",
                    level=messages.WARNING,
                )
            else:
                if resultado == "sin_correo":
                    self.message_user(
                        request,
                        "El turno se guardó, pero el cliente no tiene "
                        "un correo electrónico registrado.",
                        level=messages.WARNING,
                    )
                elif resultado == "sin_enviar":
                    self.message_user(
                        request,
                        "El servidor no confirmó el envío del correo "
                        "de avance.",
                        level=messages.WARNING,
                    )
                elif resultado == "enviado":
                    self.message_user(
                        request,
                        "Se envió al cliente el nuevo porcentaje "
                        "de reparación.",
                        level=messages.SUCCESS,
                    )
            try:
                resultado_whatsapp = enviar_whatsapp_a_cliente(
                    obj,
                    mensaje_avance_whatsapp(obj, porcentaje_anterior),
                )
            except (ImproperlyConfigured, WhatsAppDeliveryError):
                logger.exception(
                    "No se pudo enviar la actualización de WhatsApp "
                    "del turno %s desde Django Admin.",
                    obj.pk,
                )
                self.message_user(
                    request,
                    "El turno se guardó, pero no se pudo enviar "
                    "la actualización por WhatsApp.",
                    level=messages.WARNING,
                )
            else:
                if resultado_whatsapp == "enviado":
                    self.message_user(
                        request,
                        "También se envió la actualización por WhatsApp.",
                        level=messages.SUCCESS,
                    )

    def has_delete_permission(self, request, obj=None):
        return (
            request.user.is_superuser
            or request.user.has_perm("tallermami.delete_turno")
        )


# ==========================================================
# OFERTAS DE ADELANTO DE TURNOS
# ==========================================================

@admin.register(OfertaTurno)
class OfertaTurnoAdmin(admin.ModelAdmin):
    list_display = (
        "id",
        "turno_cliente",
        "fecha_disponible",
        "estado",
        "fecha_creacion",
    )

    list_filter = (
        "estado",
        "fecha_disponible",
        "fecha_creacion",
    )

    search_fields = (
        "turno_cliente__patente",
        "turno_cliente__vehiculo",
        "turno_cliente__usuario__username",
    )

    readonly_fields = ("token", "fecha_creacion")
    ordering = ("-fecha_creacion",)


@admin.register(Disponibilidad)
class DisponibilidadAdmin(admin.ModelAdmin):
    list_display = ("fecha", "hora", "estado")
    list_filter = ("fecha", "estado")
    ordering = ("fecha", "hora")


@admin.register(Vehiculo)
class VehiculoAdmin(admin.ModelAdmin):
    list_display = ("patente", "marca", "modelo", "anio", "cliente")
    search_fields = (
        "patente",
        "marca",
        "modelo",
        "cliente__nombre",
        "cliente__apellido",
        "cliente__dni",
    )
    list_filter = ("marca", "anio")
    ordering = ("patente",)


@admin.register(OrdenTrabajo)
class OrdenTrabajoAdmin(admin.ModelAdmin):
    list_display = (
        "id",
        "cliente",
        "vehiculo",
        "servicio",
        "mecanico",
        "fecha_ingreso",
        "estado",
    )
    list_filter = ("estado", "fecha_ingreso", "servicio", "mecanico")
    search_fields = (
        "cliente__nombre",
        "cliente__apellido",
        "vehiculo__patente",
        "servicio__nombre",
        "mecanico__nombre",
        "mecanico__apellido",
    )
    ordering = ("-fecha_ingreso", "-id")


@admin.register(Repuesto)
class RepuestoAdmin(admin.ModelAdmin):
    list_display = ("nombre", "stock", "precio")
    search_fields = ("nombre", "descripcion")
    ordering = ("nombre",)


@admin.register(TrabajoRealizado)
class TrabajoRealizadoAdmin(admin.ModelAdmin):
    list_display = ("id", "orden_trabajo", "descripcion", "fecha", "mecanico")
    list_filter = ("fecha", "mecanico")
    search_fields = (
        "orden_trabajo__id",
        "descripcion",
        "mecanico__nombre",
        "mecanico__apellido",
    )
    ordering = ("-fecha", "-id")


@admin.register(PerfilUsuario)
class PerfilUsuarioAdmin(admin.ModelAdmin):
    list_display = ("usuario", "rol")
    list_filter = ("rol",)
    search_fields = ("usuario__username", "usuario__email")


# ==========================================================
# REPUESTOS DE LOS PRESUPUESTOS
# ==========================================================

class RepuestoPresupuestadoInline(admin.TabularInline):
    model = RepuestoPresupuestado
    extra = 1


class RepuestoUtilizadoInline(admin.TabularInline):
    model = RepuestoUtilizado
    extra = 1


# ==========================================================
# PRESUPUESTOS
# ==========================================================

@admin.register(Presupuesto)
class PresupuestoAdmin(admin.ModelAdmin):
    list_display = (
        "id",
        "turno",
        "estado",
        "mano_obra_estimada",
        "mostrar_total_estimado",
        "mostrar_total_final",
        "final_enviado",
    )

    list_filter = ("estado", "final_enviado")

    search_fields = (
        "turno__patente",
        "turno__vehiculo",
        "turno__usuario__username",
    )

    readonly_fields = (
        "token",
        "fecha_envio",
        "fecha_respuesta",
        "fecha_finalizacion",
    )

    fields = (
        "turno",
        "estado",
        "mano_obra_estimada",
        "observaciones",
        "token",
        "fecha_envio",
        "fecha_respuesta",
        "mano_obra_final",
        "detalle_final",
        "fecha_finalizacion",
        "final_enviado",
    )

    inlines = (
        RepuestoPresupuestadoInline,
        RepuestoUtilizadoInline,
    )

    actions = (
        "enviar_presupuesto",
        "enviar_detalle_final",
    )

    @admin.display(description="Total estimado")
    def mostrar_total_estimado(self, obj):
        return f"${obj.total_estimado():,.2f}"

    @admin.display(description="Total final")
    def mostrar_total_final(self, obj):
        return f"${obj.total_final():,.2f}"

    @admin.action(description="Enviar presupuesto al cliente")
    def enviar_presupuesto(self, request, queryset):
        for presupuesto in queryset.select_related(
            "turno",
            "turno__usuario",
        ):
            turno = presupuesto.turno

            if presupuesto.estado != "pendiente":
                self.message_user(
                    request,
                    f"El presupuesto del turno #{turno.id} "
                    "ya no está pendiente.",
                    level=messages.WARNING,
                )
                continue

            email = obtener_email_cliente(turno)

            if not email:
                self.message_user(
                    request,
                    f"No se encontró un correo válido para "
                    f"el turno #{turno.id}. Revisá los datos del cliente.",
                    level=messages.ERROR,
                )
                continue

            aceptar = request.build_absolute_uri(
                reverse(
                    "responder_presupuesto",
                    args=[presupuesto.token, "aceptar"],
                )
            )

            rechazar = request.build_absolute_uri(
                reverse(
                    "responder_presupuesto",
                    args=[presupuesto.token, "rechazar"],
                )
            )

            lineas = [
                f"- {item.nombre}: {item.cantidad} x "
                f"${item.precio_unitario:.2f} = "
                f"${item.subtotal():.2f}"
                for item in presupuesto.repuestos_estimados.all()
            ]

            detalle_repuestos = (
                "\n".join(lineas)
                if lineas
                else "No se cargaron repuestos."
            )

            mensaje = (
                "Hola:\n\n"
                f"Presupuesto para el vehículo {turno.vehiculo} "
                f"(patente {turno.patente}).\n\n"
                f"Repuestos:\n{detalle_repuestos}\n\n"
                f"Mano de obra: ${presupuesto.mano_obra_estimada:.2f}\n"
                f"TOTAL ESTIMADO: ${presupuesto.total_estimado():.2f}\n\n"
                f"Observaciones: "
                f"{presupuesto.observaciones or 'Sin observaciones'}\n\n"
                "Para aceptar el presupuesto, ingresá aquí:\n"
                f"{aceptar}\n\n"
                "Para rechazarlo, ingresá aquí:\n"
                f"{rechazar}\n\n"
                "La reparación no se autorizará hasta que confirmes "
                "tu decisión.\n\n"
                "Taller Molina Henze"
            )

            try:
                enviar_correo_taller(
                    asunto="Presupuesto de reparación - Taller Molina Henze",
                    mensaje=mensaje,
                    destinatarios=[email],
                    request=request,
                    fail_silently=False,
                )

                presupuesto.fecha_envio = timezone.now()
                presupuesto.save(update_fields=["fecha_envio"])

                turno.estado = "presupuesto_pendiente"
                turno.save(update_fields=["estado"])

                self.message_user(
                    request,
                    f"Presupuesto del turno #{turno.id} enviado.",
                    level=messages.SUCCESS,
                )

            except Exception as error:
                self.message_user(
                    request,
                    f"No se pudo enviar el presupuesto del turno "
                    f"#{turno.id}: {error}",
                    level=messages.ERROR,
                )

    @admin.action(description="Enviar detalle de reparación finalizada")
    def enviar_detalle_final(self, request, queryset):
        for presupuesto in queryset.select_related(
            "turno",
            "turno__usuario",
        ):
            turno = presupuesto.turno

            if presupuesto.estado != "aceptado":
                self.message_user(
                    request,
                    f"El presupuesto del turno #{turno.id} "
                    "no está aceptado.",
                    level=messages.ERROR,
                )
                continue

            if presupuesto.final_enviado:
                self.message_user(
                    request,
                    f"El detalle final del turno #{turno.id} "
                    "ya fue enviado.",
                    level=messages.WARNING,
                )
                continue

            email = obtener_email_cliente(turno)

            if not email:
                self.message_user(
                    request,
                    f"No se encontró un correo válido para "
                    f"el turno #{turno.id}.",
                    level=messages.ERROR,
                )
                continue

            lineas = [
                f"- {item.nombre}: {item.cantidad} x "
                f"${item.precio_unitario:.2f} = "
                f"${item.subtotal():.2f}"
                for item in presupuesto.repuestos_utilizados.all()
            ]

            detalle_repuestos = (
                "\n".join(lineas)
                if lineas
                else "No se registraron repuestos utilizados."
            )

            mensaje = (
                "Hola:\n\n"
                f"Finalizamos la reparación de tu vehículo "
                f"{turno.vehiculo}, patente {turno.patente}.\n\n"
                "Porcentaje de reparación: 100%\n\n"
                f"Repuestos utilizados:\n{detalle_repuestos}\n\n"
                f"Mano de obra final: ${presupuesto.mano_obra_final:.2f}\n"
                f"TOTAL FINAL: ${presupuesto.total_final():.2f}\n\n"
                f"Detalle del trabajo: "
                f"{presupuesto.detalle_final or 'Sin detalles adicionales'}\n\n"
                "Gracias por confiar en Taller Molina Henze."
            )

            try:
                enviar_correo_taller(
                    asunto="Reparación finalizada - Taller Molina Henze",
                    mensaje=mensaje,
                    destinatarios=[email],
                    request=request,
                    fail_silently=False,
                )

                presupuesto.final_enviado = True
                presupuesto.fecha_finalizacion = timezone.now()
                presupuesto.save(
                    update_fields=[
                        "final_enviado",
                        "fecha_finalizacion",
                    ]
                )

                turno.estado = "finalizado"
                turno.porcentaje_reparacion = 100
                turno.save(
                    update_fields=[
                        "estado",
                        "porcentaje_reparacion",
                    ]
                )

                self.message_user(
                    request,
                    f"Detalle final del turno #{turno.id} enviado.",
                    level=messages.SUCCESS,
                )

            except Exception as error:
                self.message_user(
                    request,
                    f"No se pudo enviar el detalle final del turno "
                    f"#{turno.id}: {error}",
                    level=messages.ERROR,
                )


# ==========================================================
# REPUESTOS: ADMINISTRACIÓN INDEPENDIENTE
# ==========================================================

@admin.register(RepuestoPresupuestado)
class RepuestoPresupuestadoAdmin(admin.ModelAdmin):
    list_display = (
        "nombre",
        "presupuesto",
        "cantidad",
        "precio_unitario",
    )
    search_fields = ("nombre",)


@admin.register(RepuestoUtilizado)
class RepuestoUtilizadoAdmin(admin.ModelAdmin):
    list_display = (
        "nombre",
        "presupuesto",
        "cantidad",
        "precio_unitario",
    )
    search_fields = ("nombre",)
